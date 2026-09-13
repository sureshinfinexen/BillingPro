"""Generate BillingPro complete user operations Word document."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

OUT = Path(r"D:\Source\BillingPro\BillingPro_Complete_User_Operations_Guide.docx")


def H(doc, text, level=1):
    return doc.add_heading(text, level=level)


def P(doc, text, bold=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(11)
    return p


def bullets(doc, items):
    for i in items:
        doc.add_paragraph(i, style="List Bullet")


def numbered(doc, items):
    for i in items:
        doc.add_paragraph(i, style="List Number")


def table(doc, headers, rows):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = h
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            t.rows[ri + 1].cells[ci].text = str(val)
    doc.add_paragraph()
    return t


def note(doc, text):
    p = doc.add_paragraph()
    r = p.add_run("Note: ")
    r.bold = True
    r.font.color.rgb = RGBColor(0x0F, 0x76, 0x6E)
    r2 = p.add_run(text)
    r2.font.size = Pt(11)


def build():
    doc = Document()

    # ── Title ──
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("BillingPro")
    r.bold = True
    r.font.size = Pt(28)
    r.font.color.rgb = RGBColor(0x0F, 0x76, 0x6E)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run("Complete User Operations Guide")
    r.bold = True
    r.font.size = Pt(18)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = meta.add_run(
        "Point of Sale & Billing  ·  Infinexen Technologies  ·  2026\n"
        "English / Tamil / Hindi  ·  Version aligned with current application"
    )
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(0x5D, 0x6D, 0x7E)

    P(doc, "")
    P(
        doc,
        "This guide explains how to set up, sell, and operate BillingPro day to day — "
        "including Super Admin feature packs, store admin page assignment, and cashier billing.",
    )

    # ═══════════════════════════════════════════════════════════
    H(doc, "1. What is BillingPro?")
    P(
        doc,
        "BillingPro is a web-based Point of Sale (POS) and billing system for retail shops. "
        "It supports GST billing, stock, multi-counter cashiers, holds/returns, credit (udhaar), "
        "loyalty, coupons, purchases, warehouses, day-close, GST export, and multi-shop licensing.",
    )
    bullets(
        doc,
        [
            "Open in a browser: http://localhost:8003 (default port).",
            "Start with START_BILLINGPRO.bat or: python run.py",
            "Full install steps: docs/SETUP.md",
        ],
    )

    # ═══════════════════════════════════════════════════════════
    H(doc, "2. How to log in")
    numbered(
        doc,
        [
            "Open the application URL in Chrome / Edge.",
            "Enter Username and Password.",
            "Click Login.",
            "Use the language buttons (EN / TA / HI) in the top bar anytime after login.",
        ],
    )
    H(doc, "2.1 Default accounts (change in production)", level=2)
    table(
        doc,
        ["Role", "Username", "Password", "After login goes to"],
        [
            ["Super Admin", "superadmin", "superadmin123", "Shops (platform)"],
            ["Store Admin", "admin", "admin123", "Dashboard"],
            ["Cashier", "cashier", "cashier123", "New Bill"],
        ],
    )
    note(doc, "Change all default passwords before giving the system to a customer store.")

    # ═══════════════════════════════════════════════════════════
    H(doc, "3. Roles & access model")
    P(
        doc,
        "Every screen is allowed only if: (1) the user role allows it, "
        "(2) the page is assigned to that user, and (3) the shop feature pack enables that module. "
        "Super Admin bypasses all checks.",
    )
    P(doc, "Hierarchy", bold=True)
    P(
        doc,
        "Super Admin → enables feature packs per Shop + creates Store Admin\n"
        "    → Store Admin → assigns pages to Cashiers\n"
        "        → Cashier → uses only assigned + enabled screens",
    )
    table(
        doc,
        ["Role", "Scope", "Main responsibilities"],
        [
            [
                "Super Admin",
                "All shops (platform)",
                "Create shops, enable/disable sold modules, create store admins, full access",
            ],
            [
                "Store Admin (admin)",
                "One shop",
                "Products, stock, employees, counters, reports, Assign pages",
            ],
            [
                "Cashier",
                "One shop + counter",
                "Billing and other screens assigned by store admin",
            ],
        ],
    )

    # ═══════════════════════════════════════════════════════════
    H(doc, "4. Left-side menus (navigation)")
    P(
        doc,
        "Menus are grouped. Items appear only when allowed by role + assigned pages + shop features.",
    )
    table(
        doc,
        ["Group", "Menu items"],
        [
            ["Overview", "Dashboard, Reports, Shifts, Day close"],
            ["Billing", "New Bill, Held bills, Bills, Returns, Credit (Udhaar)"],
            ["Inventory", "Stock, Reorder, Products, Purchases, Warehouses, Batches"],
            ["Customers", "Loyalty, Coupons"],
            ["Reports & tax", "GST export, Audit log"],
            [
                "Administration",
                "Shops, Counters, Employees, Assign pages, Offline, Settings",
            ],
        ],
    )

    # ═══════════════════════════════════════════════════════════
    H(doc, "5. Correct setup order (do this first)")
    P(
        doc,
        "Follow this order when installing for a new store or when selling BillingPro to another shop.",
    )
    numbered(
        doc,
        [
            "Install app + PostgreSQL database (see docs/SETUP.md).",
            "Login as Super Admin.",
            "Create the Shop (name, code, address) and tick Feature pack modules sold to that store.",
            "Create a Store Admin for that shop (Features & admin → Create store admin).",
            "Login as Store Admin → Settings (app name, theme, default language, bill greeting).",
            "Create Counters (e.g. Counter 1, Counter 2).",
            "Add Products (manually or CSV upload) with MRP, discount, rate, cost, GST%, stock, HSN.",
            "Create Employees (cashiers) and assign each to a Counter.",
            "Open Assign pages and tick screens each cashier may use.",
            "Train cashiers: New Bill → pay → print. Optional: holds, returns, day close.",
            "Ask each user to log out and log in again after page/feature changes.",
        ],
    )

    # ═══════════════════════════════════════════════════════════
    H(doc, "6. Super Admin — shops & feature packs (selling)")
    H(doc, "6.1 Add a shop", level=2)
    numbered(
        doc,
        [
            "Login as superadmin.",
            "Go to Administration → Shops / branches.",
            "Fill Name, Code, Address.",
            "Tick Feature pack modules the customer paid for.",
            "Click Save.",
        ],
    )
    H(doc, "6.2 Manage features & create store admin", level=2)
    numbered(
        doc,
        [
            "In the shops list, click Features & admin for that shop.",
            "Enable / disable modules and Save.",
            "Under Create store admin, enter Name, Username, Password → Save.",
            "Give those credentials to the store owner/manager.",
        ],
    )
    H(doc, "6.3 Feature pack catalogue", level=2)
    table(
        doc,
        ["Feature", "What it unlocks", "Core?"],
        [
            ["Billing", "New Bill, Bills list / view / edit", "Yes"],
            ["Stock", "Stock view & adjust", "Yes"],
            ["Products", "Product master + CSV", "Yes"],
            ["Holds", "Hold / recall bills", "No"],
            ["Returns", "Returns / credit notes", "No"],
            ["Credit", "Credit sales (Udhaar) + payments", "No"],
            ["Loyalty", "Loyalty points", "No"],
            ["Coupons", "Coupons / promotions", "No"],
            ["Reorder", "Low-stock reorder list", "No"],
            ["Purchases", "Purchases / GRN", "No"],
            ["Warehouses", "Warehouses & stock transfers", "No"],
            ["Batches", "Batch / expiry tracking", "No"],
            ["Day close", "Day-end closing", "No"],
            ["Shifts", "Shift / counter sales reports", "No"],
            ["Reports", "Dashboard + sales reports", "No"],
            ["GST export", "GST CSV export", "No"],
            ["Audit", "Audit log", "No"],
            ["Offline", "Offline mode helper", "No"],
        ],
    )
    note(doc, "Core modules stay on for every licensed shop. Optional modules are what you sell as packs.")

    # ═══════════════════════════════════════════════════════════
    H(doc, "7. Store Admin — Assign pages to users")
    numbered(
        doc,
        [
            "Login as store admin.",
            "Go to Administration → Assign pages.",
            "For each employee, tick the screens they may open.",
            "Only pages whose shop feature is enabled are listed.",
            "Click Save for that user.",
            "Ask the user to log out and log in again.",
        ],
    )
    P(doc, "Typical cashier page set", bold=True)
    bullets(
        doc,
        [
            "New Bill, Bills, Held bills, Stock",
            "Optional: Returns, Credit, Day close, Reorder",
        ],
    )
    P(doc, "Default cashier pages if none saved yet: billing, bills, holds, stock, returns, credit, day_close, reorder.")

    # ═══════════════════════════════════════════════════════════
    H(doc, "8. Counters")
    P(doc, "Counters represent physical billing desks (Counter 1, Counter 2, …).")
    numbered(
        doc,
        [
            "Administration → Counters.",
            "Enter Name, Code (e.g. C1), Location.",
            "Save. Assign cashiers to a counter under Employees.",
        ],
    )

    # ═══════════════════════════════════════════════════════════
    H(doc, "9. Employees (staff)")
    numbered(
        doc,
        [
            "Administration → Employees.",
            "Enter Name, Username, Password.",
            "Choose Role: admin (store admin) or cashier.",
            "Select Counter for cashiers.",
            "Save.",
            "Then open Assign pages for that user.",
        ],
    )
    note(doc, "You cannot create a Super Admin from the Employees screen. Only Super Admin creates store admins from Shops → Features & admin.")

    # ═══════════════════════════════════════════════════════════
    H(doc, "10. Products")
    H(doc, "10.1 Add / edit a product", level=2)
    numbered(
        doc,
        [
            "Inventory → Products → Add (or Edit).",
            "Enter SKU (unique barcode/code), Name.",
            "Enter MRP, Discount (₹), Rate is calculated as MRP − Discount.",
            "Enter Cost, Stock, Unit, Category, GST %, HSN, Reorder level.",
            "Save.",
        ],
    )
    H(doc, "10.2 CSV upload", level=2)
    P(
        doc,
        "Upload a CSV with columns such as: sku, name, mrp, discount, price/rate, cost, stock, unit, category, gst/gst_pct. "
        "Existing SKUs are updated; new SKUs are inserted.",
    )

    # ═══════════════════════════════════════════════════════════
    H(doc, "11. Stock")
    numbered(
        doc,
        [
            "Inventory → Stock.",
            "Search by name or SKU.",
            "Adjust quantity: Set / Add / Remove.",
            "Creating a bill deducts stock; editing a bill restores old lines then deducts new lines.",
        ],
    )
    P(doc, "Reorder list (if enabled) shows products at or below reorder level.")

    # ═══════════════════════════════════════════════════════════
    H(doc, "12. New Bill (POS) — daily cashier workflow")
    numbered(
        doc,
        [
            "Open New Bill (cashiers land here after login).",
            "Confirm Counter (cashiers are usually fixed to one counter).",
            "Search product by name or SKU, or scan barcode (exact SKU).",
            "Set Qty. Check MRP, Discount, Rate, GST.",
            "Optional: customer name / mobile (for credit, loyalty, WhatsApp).",
            "Optional: coupon code, loyalty redeem, bill-level discount (if enabled).",
            "Choose Payment: Cash / UPI / Card. For Card, enter last 4 digits (mandatory).",
            "Optional: Credit (Udhaar) if enabled and customer identified.",
            "Save / Complete bill. Print or open thermal reprint / WhatsApp link from bill view.",
        ],
    )
    H(doc, "12.1 Hold / Recall", level=2)
    bullets(
        doc,
        [
            "Hold saves the current cart so you can serve another customer.",
            "Held bills → Recall to continue, or delete the hold.",
        ],
    )
    H(doc, "12.2 Bills list", level=2)
    bullets(
        doc,
        [
            "View recent bills, open bill detail, print, reprint, WhatsApp share.",
            "Edit bill (if permitted): stock is corrected automatically.",
            "Cashiers typically see their own bills; admins see all.",
        ],
    )

    # ═══════════════════════════════════════════════════════════
    H(doc, "13. Returns")
    numbered(
        doc,
        [
            "Billing → Returns.",
            "Enter original Bill No and load items.",
            "Select quantities to return and reason.",
            "Save return — stock is added back as configured.",
        ],
    )

    # ═══════════════════════════════════════════════════════════
    H(doc, "14. Credit sales (Udhaar)")
    numbered(
        doc,
        [
            "Create a bill with credit / udhaar payment (feature must be enabled).",
            "Open Credit screen to see outstanding balances.",
            "Record customer payment against outstanding credit.",
        ],
    )

    # ═══════════════════════════════════════════════════════════
    H(doc, "15. Loyalty & Coupons")
    bullets(
        doc,
        [
            "Loyalty: points earn/redeem rules from settings; manage/view under Customers → Loyalty.",
            "Coupons: create codes with % or amount off and minimum bill (sample seed: SAVE10).",
            "Apply coupon on the billing screen when the feature is enabled.",
        ],
    )

    # ═══════════════════════════════════════════════════════════
    H(doc, "16. Purchases, Warehouses, Batches")
    bullets(
        doc,
        [
            "Purchases / GRN — receive supplier stock into inventory.",
            "Warehouses — multiple godowns; transfer stock between warehouses.",
            "Batches — track batch number and expiry where required.",
        ],
    )

    # ═══════════════════════════════════════════════════════════
    H(doc, "17. Day close & Shifts")
    H(doc, "17.1 Day close", level=2)
    numbered(
        doc,
        [
            "Open Day close.",
            "Review expected Cash / UPI / Card from today’s bills.",
            "Enter actual cash counted and notes.",
            "Save closing for the day / counter.",
        ],
    )
    H(doc, "17.2 Shifts", level=2)
    P(doc, "Shift / counter sales report shows sales broken down by counter and staff for monitoring.")

    # ═══════════════════════════════════════════════════════════
    H(doc, "18. Dashboard")
    bullets(
        doc,
        [
            "Today’s sales total and bill count.",
            "Active product count.",
            "Recent bills list for quick review.",
        ],
    )

    # ═══════════════════════════════════════════════════════════
    H(doc, "19. Reports & GST export")
    bullets(
        doc,
        [
            "Reports — revenue, discount, tax, profit, best buyer, best product.",
            "GST export — download CSV for a date range (GSTR-style line details).",
            "Audit log — who did what (permissions changes, reprints, feature updates, etc.).",
        ],
    )

    # ═══════════════════════════════════════════════════════════
    H(doc, "20. Settings")
    bullets(
        doc,
        [
            "App name and tagline.",
            "Default language (English / Tamil / Hindi).",
            "Theme presets (light/dark styles) and custom colours.",
            "Bill greeting and thank-you lines printed on bills.",
            "Reset theme to default when needed.",
        ],
    )

    # ═══════════════════════════════════════════════════════════
    H(doc, "21. Offline helper")
    P(
        doc,
        "Offline mode page helps cache basic offline behaviour and sync queued work when back online. "
        "Enable the Offline feature in the shop pack if you sell this module.",
    )

    # ═══════════════════════════════════════════════════════════
    H(doc, "22. How modules link together")
    P(
        doc,
        "Shop (feature pack)\n"
        "  ├── Counters\n"
        "  ├── Employees (pages assigned) ──► New Bill / Stock / …\n"
        "  ├── Products (MRP, GST, stock)\n"
        "  │     ├── Stock adjustments / Purchases / Batches\n"
        "  │     └── Bill lines (qty deducted)\n"
        "  ├── Customers (mobile) ──► Credit, Loyalty, WhatsApp bill\n"
        "  ├── Coupons / Holds / Returns\n"
        "  └── Day close / Shifts / Reports / GST export / Audit",
    )

    # ═══════════════════════════════════════════════════════════
    H(doc, "23. Daily & weekly workflows")
    H(doc, "23.1 Cashier — daily", level=2)
    numbered(
        doc,
        [
            "Login → New Bill.",
            "Bill customers throughout the day; use Hold when needed.",
            "Process returns / credit collections if assigned.",
            "Day close at end of shift (if assigned).",
            "Logout.",
        ],
    )
    H(doc, "23.2 Store admin — daily / weekly", level=2)
    numbered(
        doc,
        [
            "Check Dashboard and Reports.",
            "Review low stock / Reorder; receive Purchases.",
            "Add new products; adjust stock if required.",
            "Review Shift reports and Audit if needed.",
            "Export GST CSV for accountant (weekly/monthly).",
            "Add cashiers and update Assign pages when staff change.",
        ],
    )
    H(doc, "23.3 Super Admin — when selling / renewing", level=2)
    numbered(
        doc,
        [
            "Create or select the shop.",
            "Enable only paid feature modules.",
            "Create / reset store admin login.",
            "Confirm store admin can Assign pages and run POS.",
        ],
    )

    # ═══════════════════════════════════════════════════════════
    H(doc, "24. Troubleshooting")
    table(
        doc,
        ["Problem", "What to do"],
        [
            ["Cannot connect / blank page", "Check PostgreSQL and POSTGRES_DSN in .env; restart app"],
            ["Login fails", "Confirm username active; reset password via Employees (store admin)"],
            ["Menu item missing", "Check shop Features (superadmin) + Assign pages + re-login"],
            ["Dashboard error after update", "Hard refresh (Ctrl+F5); restart BillingPro; re-login"],
            ["Card payment blocked", "Enter last 4 digits of card"],
            ["Stock wrong after edit", "Edit bill again carefully; stock restore/deduct runs in transaction"],
            ["Port in use", "Change PORT in .env (default 8003)"],
        ],
    )

    # ═══════════════════════════════════════════════════════════
    H(doc, "25. Related documents")
    bullets(
        doc,
        [
            "docs/SETUP.md — install from zero (Python, Postgres, .env, start).",
            "docs/FUNCTIONALITY.md — technical feature & access reference.",
            "README.md — short project overview.",
        ],
    )

    # ═══════════════════════════════════════════════════════════
    H(doc, "26. Contact")
    P(doc, "Infinexen Technologies")
    bullets(
        doc,
        [
            "Email: suresh.baskaran@infinexen.com",
            "Mobile: +91 87780 99759",
            "Year: 2026",
        ],
    )

    footer = doc.add_paragraph()
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run("— End of BillingPro Complete User Operations Guide —")
    r.italic = True
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(0x5D, 0x6D, 0x7E)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    build()
